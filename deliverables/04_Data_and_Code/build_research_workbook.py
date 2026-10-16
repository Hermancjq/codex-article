#!/usr/bin/env python3
"""Package frozen research outputs as Excel; never rerun the analysis."""
from pathlib import Path
import csv
import datetime as dt
import hashlib
import json
import math
import re
from urllib.parse import urlsplit

import openpyxl
import xlsxwriter

ROOT = Path('/workspace/research')
OUT = ROOT / 'delivery/04_Data_and_Code'
DEST = OUT / 'RESEARCH_DATA_AND_RESULTS.xlsx'
CSV_INPUTS = [
    ('Verified_Observations', 'source/verified_data.csv', 'Official observations / documented scope interpretation'),
    ('Raw_Official_2010_2025', 'data_agent/official_aviation_2010_2025.csv', 'Official observation ledger; early fuel intensity missing'),
    ('NBS_Check_2025_Yearbook', 'data_agent/nbs_crosscheck.csv', 'Official cross-publication check'),
    ('NBS_Cross_Publication', 'scope_audit/nbs_turnover_crosscheck.csv', 'Official cross-publication check; not independent measurement'),
    ('Annual_Reconstruction', 'output/analysis/annual.csv', 'Official Q and I plus reconstructed fuel and indices'),
    ('LMDI_Annual', 'output/analysis/decomposition.csv', 'Accounting results calculated from Q and I'),
    ('LMDI_Periods', 'output/analysis/periods.csv', 'Frozen period decomposition results'),
    ('Forecast_Diagnostic', 'output/transition/rolling_h1_forecasts.csv', 'Small-sample diagnostic forecasts'),
    ('Scenario_Design', 'output/transition/conditional_scenario_design.csv', 'Conditional stress-test design; not probabilities'),
]
JSON_INPUTS = ['scenario_config.json', 'output/transition/transition_results.json']
ALL_INPUTS = [p for _, p, _ in CSV_INPUTS] + JSON_INPUTS
SHA_BEFORE = {p: hashlib.sha256((ROOT / p).read_bytes()).hexdigest() for p in ALL_INPUTS}
NUMERIC = re.compile(r'[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?\Z')
INTEGER = re.compile(r'[+-]?\d+\Z')
csv_data = {}
for name, path, category in CSV_INPUTS:
    with (ROOT / path).open(encoding='utf-8-sig', newline='') as f:
        rows = list(csv.reader(f))
    assert rows and all(len(r) == len(rows[0]) for r in rows), path
    csv_data[name] = (path, category, rows[0], rows[1:])
config = json.loads((ROOT / JSON_INPUTS[0]).read_text())
transition = json.loads((ROOT / JSON_INPUTS[1]).read_text())

workbook = xlsxwriter.Workbook(str(DEST), {'strings_to_formulas': False, 'strings_to_urls': False})
workbook.set_properties({'title': 'Official Aviation Data and Frozen Research Results',
                         'subject': 'Observed data, accounting estimates, and conditional deployment scenarios',
                         'author': 'Research delivery',
                         'comments': 'Packaging only. No model, data extraction, or statistical analysis was rerun.'})
header_fmt = workbook.add_format({'bold': True, 'font_color': '#FFFFFF', 'bg_color': '#17324D',
                                  'text_wrap': True, 'valign': 'vcenter', 'border': 0, 'font_size': 10})
format_cache = {}
def style(number_format=None, alternate=False, url=False):
    key = (number_format, alternate, url)
    if key not in format_cache:
        opts = {'font_name': 'Aptos', 'font_size': 10, 'valign': 'top', 'text_wrap': True,
                'bg_color': '#F0F5FA' if alternate else '#FFFFFF'}
        if number_format:
            opts['num_format'] = number_format
        if url:
            opts.update({'font_color': '#0563C1', 'underline': 1})
        format_cache[key] = workbook.add_format(opts)
    return format_cache[key]

sheet_specs = {}
def number_format(header, value):
    h = header.lower()
    if any(w in h for w in ['year', 'seed', 'sobol_power', 'n_design_points', 'n_valid', 'horizon', 'origin', 'start', 'end']) and isinstance(value, int):
        return '0'
    if h == 'year' or h == 'observation_year' or h in ['tau', 'start', 'end', 'horizon', 'n_design_points', 'seed', 'sobol_power']:
        return '0'
    if 'residual' in h:
        return '0.000E+00'
    if 'intensity_offset_of_activity_percent' == h:
        return '0.00"%"'
    if any(w in h for w in ['growth', 'cagr', 'fraction_of_design_points', 'break_even_intensity_reduction']):
        return '0.00%'
    if h.startswith('fuel_intensity'):
        return '0.000'
    if 'difference' in h:
        return '0.000000'
    if h.endswith('10000_tkm'):
        return '#,##0'
    if h in ['turnover_100m_tkm', 'caac_bulletin_turnover_100m_tkm', 'caac_bulletin_value_100m_tkm']:
        return '#,##0.00'
    if '100m_tkm' in h:
        return '#,##0.0000'
    if any(w in h for w in ['ratio', 'index', '10kt', '_mt']):
        return '#,##0.000000'
    return '#,##0' if isinstance(value, int) else '0.0000000000'

def is_url(s):
    p = urlsplit(s)
    return p.scheme in ('https', 'http') and bool(p.hostname) and not p.username and not p.password

def write_sheet(name, headers, rows, numeric_columns=None, long_rows=False):
    ws = workbook.add_worksheet(name)
    ws.freeze_panes(1, 0)
    ws.hide_gridlines(2)
    ws.set_zoom(90)
    ws.set_row(0, 34)
    ws.set_default_row(21)
    ws.set_landscape()
    ws.fit_to_pages(1, 0)
    ws.repeat_rows(0)
    ws.set_header('&L&BOFFICIAL DATA & RESEARCH RESULTS&R' + name)
    ws.set_footer('&LSource values frozen; no recalculation&CPage &P of &N')
    if numeric_columns is None:
        numeric_columns = set()
    for col, header in enumerate(headers):
        ws.write_string(0, col, header, header_fmt)
        lower = header.lower()
        if 'sha256' in lower:
            width = 68
        elif 'url' in lower or 'source_file' in lower:
            width = 56
        elif any(t in lower for t in ['quote', 'coverage', 'scope_notes', 'description', 'interpretation', 'literal', 'meaning', 'independence', 'value_json']):
            width = 78
        elif 'locator' in lower or 'page_or_table' in lower or 'transcription' in lower:
            width = 48
        elif len(header) > 26:
            width = 28
        else:
            width = max(14, min(25, len(header) + 3))
        ws.set_column(col, col, width)
    for ridx, row in enumerate(rows, 1):
        if long_rows:
            ws.set_row(ridx, 68)
        for cidx, raw in enumerate(row):
            alternate = bool(ridx % 2 == 0)
            header = headers[cidx]
            if raw is None or raw == '':
                ws.write_blank(ridx, cidx, None, style(alternate=alternate))
            elif cidx in numeric_columns:
                value = int(raw) if isinstance(raw, str) and INTEGER.fullmatch(raw) else float(raw)
                assert math.isfinite(value), (name, ridx, cidx)
                ws.write_number(ridx, cidx, value, style(number_format(header, value), alternate))
            elif isinstance(raw, bool):
                ws.write_boolean(ridx, cidx, raw, style(alternate=alternate))
            elif isinstance(raw, (int, float)):
                assert math.isfinite(raw)
                ws.write_number(ridx, cidx, raw, style(number_format(header, raw), alternate))
            else:
                value = str(raw)
                assert len(value) <= 32767, (name, ridx, cidx, len(value))
                if is_url(value):
                    ws.write_url(ridx, cidx, value, style(alternate=alternate, url=True), string=value)
                else:
                    ws.write_string(ridx, cidx, value, style(alternate=alternate))
    ws.autofilter(0, 0, len(rows), len(headers) - 1)
    ws.print_area(0, 0, len(rows), len(headers) - 1)
    sheet_specs[name] = {'headers': headers, 'rows': rows, 'numeric_columns': numeric_columns}
    return ws

readme_rows = [
    ['Purpose / 用途', 'Downloadable Excel packaging of frozen official-source observations, accounting results and explicitly conditional stress tests. No input, analysis script, manuscript or numerical result was changed; no model was rerun.'],
    ['Primary window / 主分析窗', '2017–2025. The 2017 CAAC bulletin reports changed production-data processing and audit rules. Retain 2015–2016 as an extended-window sensitivity, not an undisclosed part of the primary window.'],
    ['Observed data / 官方观测', 'Verified_Observations preserves 2015–2025 annual Q and I with official URLs, PDF locators, full SHA-256 and the documented population interpretation. Raw_Official_2010_2025 preserves 16 years of Q; I is missing for 2010–2014 and remains blank.'],
    ['Scope / 统计口径', 'Q covers mainland-registered commercial carriers on domestic and international routes. I is interpreted as the transport-aviation series using official anchors and series continuity, not annual matched carrier microdata. Read the original coverage_definition and scope_notes cells.'],
    ['NBS cross-check / 年鉴核对', 'NBS sheets compare publications and rounding. The NBS yearbooks reuse CAAC-source statistics: this is cross-publication verification, not an independent original measurement of traffic. Both original check files are preserved, including overlapping years.'],
    ['Reconstructed results / 推算结果', 'Annual_Reconstruction and LMDI sheets retain the existing Q/I accounting outputs. Implied fuel is reconstructed as Q × I with unit conversion; it is not an independently observed official fuel total. No independent aggregate fuel-total validation is claimed.'],
    ['Units / 单位', 'Q: 100 million tonne-km (亿吨公里); I: kg fuel per tonne-km; implied_fuel_mt: million tonnes; *_10kt: 10,000 tonnes. The annual index columns retain the frozen 2015=100 base; they have not been rebased for the 2017 primary window.'],
    ['Rounding / 原始精度', 'Reported I is rounded to 0.001 kg/tkm. Numeric display formats only affect appearance. Full original CSV tokens are retained visibly in CSV_Literal_Archive to preserve all source digits beyond Excel numeric precision.'],
    ['Forecast diagnostic / 预测诊断', 'Forecast_Diagnostic contains only the frozen one-year-ahead diagnostic forecasts. There are two test origins (2023 and 2024), so the table does not establish reliable model superiority. It is separate from the conditional deployment scenarios.'],
    ['Conditional scenarios / 条件压力情景', 'Scenario_Design and Delay_Summaries_HighRes are analyst-defined stress tests, not calibrated real-world distributions. Fractions of finite design points meeting a target are NOT target probabilities; design quantiles are NOT confidence intervals.'],
    ['Scenario time / 情景时间', 'Scenario_Config retains the original 15-year horizon: observed 2025 baseline to 2040. The terminal fuel ratio target of 1.0 is an analytical benchmark, not an official policy target or a claim of climate neutrality.'],
    ['Scenario parameters / 参数', 'g and b are continuous log rates per year. theta is an additional realized efficiency advantage relative to the improving baseline, not an engine-only gain. r is fuel-use-weighted adoption, not an aircraft-count replacement rate. tau is deployment delay in years.'],
    ['Design resolutions / 设计分辨率', 'Scenario_Design contains all 12,288 frozen CSV rows: 4,096 points for each delay. Delay_Summaries_HighRes copies the already computed 16,384-point summaries for each delay from transition_results.json; no high-resolution scenarios were rerun.'],
    ['Inference limits / 推断边界', 'No domestic-engine R&D capability, engine-specific performance, causal technology effect, lifecycle CO2-equivalent, or total aviation climate impact is identified by these aggregates or scenarios.'],
    ['Literal storage / 文本安全', 'All source strings are written as literal Excel strings. Auto formula conversion and auto URL conversion are disabled. Only complete ordinary HTTP/HTTPS URLs receive explicit hyperlinks. No worksheet contains formulas or hidden rows.'],
    ['Audit / 核验', 'Source_Index records each frozen input path and SHA-256. CSV_Literal_Archive preserves every CSV data row as exact-token JSON with its source row number. JSON_Literal_Archive preserves original JSON text in ordered chunks. excel_validation.json checks rows, headers, cell values, source hashes, formula absence and visibility.'],
]
for name, path, category in CSV_INPUTS:
    readme_rows.append([name, category + ' | Frozen source: ' + path])
write_sheet('README', ['Topic', 'Description'], readme_rows, long_rows=True)

literal_rows = []
for name, path, category in CSV_INPUTS:
    _, _, headers, rows = csv_data[name]
    numeric_cols = {i for i in range(len(headers))
                    if any(r[i] != '' for r in rows) and all(r[i] == '' or NUMERIC.fullmatch(r[i]) for r in rows)
                    and not any(t in headers[i].lower() for t in ['sha256', 'url', 'locator', 'source_file'])}
    ws = write_sheet(name, headers, rows, numeric_cols,
                     long_rows=name in ['Verified_Observations', 'Raw_Official_2010_2025', 'NBS_Cross_Publication'])
    for i, h in enumerate(headers):
        if name == 'Scenario_Design' and h in ['g', 'b', 'theta', 'r', 'tau', 'meets_target']:
            ws.write_comment(0, i, 'Conditional stress-test parameter or finite-design result; not a calibrated probability. See README and Scenario_Config.')
        if 'source_sha256' in h:
            ws.write_comment(0, i, 'Complete SHA-256 of the archived source bytes, retained as literal text.')
    for row_number, row in enumerate(rows, 2):
        literal_rows.append([path, row_number, json.dumps(dict(zip(headers, row)), ensure_ascii=False, separators=(',', ':'))])

def flatten(value, prefix=''):
    if isinstance(value, dict):
        for k, v in value.items():
            yield from flatten(v, prefix + '.' + k if prefix else k)
    elif isinstance(value, list):
        if not value:
            yield [prefix, '[]', 'empty_list']
        for idx, v in enumerate(value):
            yield from flatten(v, f'{prefix}[{idx}]')
    else:
        yield [prefix, value, type(value).__name__]

config_rows = list(flatten(config))
write_sheet('Scenario_Config', ['parameter_path', 'value', 'value_type'], config_rows)
scenario = transition['conditional_deployment_scenarios']
delay_headers = ['tau', 'sobol_power', 'n_design_points', 'seed', 'fraction_of_design_points_meeting_target',
                 'terminal_ratio_design_q05', 'terminal_ratio_design_q50', 'terminal_ratio_design_q95',
                 'cumulative_design_q05_baseline_years', 'cumulative_design_q50_baseline_years', 'cumulative_design_q95_baseline_years',
                 'sampled_terminal_min', 'sampled_terminal_max', 'change_in_design_target_fraction',
                 'change_in_terminal_median', 'source_json_pointer']
def delay_row(d, pointer, changes=None):
    t = d['terminal_fuel_ratio_design_quantiles']
    c = d['cumulative_fuel_design_quantiles_baseline_year_units']
    return [d['tau'], d['sobol_power'], d['n_design_points'], d['seed'], d['fraction_of_design_points_meeting_target'],
            t['0.05'], t['0.5'], t['0.95'], c['0.05'], c['0.5'], c['0.95'],
            d['sampled_terminal_min'], d['sampled_terminal_max'],
            None if changes is None else changes['change_in_design_target_fraction'],
            None if changes is None else changes['change_in_terminal_median'], pointer]
high_rows = [delay_row(x['higher_resolution'], f'/conditional_deployment_scenarios/resolution_comparison/{i}/higher_resolution', x)
             for i, x in enumerate(scenario['resolution_comparison'])]
write_sheet('Delay_Summaries_HighRes', delay_headers, high_rows)
base_rows = [delay_row(x, f'/conditional_deployment_scenarios/delay_summaries/{i}')
             for i, x in enumerate(scenario['delay_summaries'])]
write_sheet('Delay_Summaries_Base', delay_headers, base_rows)
bound_headers = ['corner', 'g', 'b', 'theta', 'r', 'tau', 'terminal_fuel_ratio',
                 'cumulative_future_fuel_in_baseline_year_units', 'source_json_pointer']
bound_rows = []
for key in ['best_corner', 'worst_corner']:
    item = scenario['exact_domain_bounds'][key]
    bound_rows.append([key] + [item['parameters'][k] for k in ['g', 'b', 'theta', 'r', 'tau']] +
                      [item['terminal_fuel_ratio'], item['cumulative_future_fuel_in_baseline_year_units'],
                       '/conditional_deployment_scenarios/exact_domain_bounds/' + key])
write_sheet('Exact_Domain_Bounds', bound_headers, bound_rows)
source_rows = [[name, path, len(csv_data[name][3]), len(csv_data[name][2]), SHA_BEFORE[path], category]
               for name, path, category in CSV_INPUTS]
source_rows += [['Scenario_Config', JSON_INPUTS[0], len(config_rows), 3, SHA_BEFORE[JSON_INPUTS[0]], 'Original scenario assumptions, flattened without numerical changes'],
                ['Delay_Summaries_HighRes', JSON_INPUTS[1], len(high_rows), len(delay_headers), SHA_BEFORE[JSON_INPUTS[1]], 'Higher-resolution summaries already present in frozen JSON']]
write_sheet('Source_Index', ['sheet', 'frozen_input_relative_path', 'data_rows', 'columns', 'input_sha256', 'classification'], source_rows)
write_sheet('CSV_Literal_Archive', ['source_csv', 'source_csv_row_number', 'exact_source_values_json'], literal_rows)
json_literal_rows = []
for path in JSON_INPUTS:
    text = (ROOT / path).read_text(encoding='utf-8')
    chunks = [text[i:i + 30000] for i in range(0, len(text), 30000)]
    for idx, chunk in enumerate(chunks, 1):
        json_literal_rows.append([path, idx, len(chunks), chunk])
write_sheet('JSON_Literal_Archive', ['source_json', 'chunk_ordinal', 'total_chunks', 'literal_original_json_chunk'], json_literal_rows)
workbook.close()

# Independent workbook readback: do not run any statistical or simulation code.
book = openpyxl.load_workbook(DEST, data_only=False, read_only=False)
errors = []
report_sheets = []
formula_cells = []
for name, spec in sheet_specs.items():
    ws = book[name]
    rows, headers, numeric_cols = spec['rows'], spec['headers'], spec['numeric_columns']
    if ws.max_row != len(rows) + 1 or ws.max_column != len(headers):
        errors.append(f'{name}: shape mismatch')
    if [ws.cell(1, i + 1).value for i in range(len(headers))] != headers:
        errors.append(f'{name}: header mismatch')
    if ws.freeze_panes != 'A2' or not ws.auto_filter.ref:
        errors.append(f'{name}: missing freeze/filter')
    if ws.sheet_state != 'visible' or any(d.hidden for d in ws.row_dimensions.values()) or any(d.hidden for d in ws.column_dimensions.values()):
        errors.append(f'{name}: hidden data')
    for ridx, row in enumerate(rows, 2):
        for cidx, original in enumerate(row, 1):
            cell = ws.cell(ridx, cidx)
            if cell.data_type == 'f':
                formula_cells.append(f'{name}!{cell.coordinate}')
            actual = cell.value
            if original is None or original == '':
                matches = actual is None
            elif cidx - 1 in numeric_cols or isinstance(original, (int, float)) and not isinstance(original, bool):
                matches = isinstance(actual, (int, float)) and not isinstance(actual, bool) and math.isclose(float(actual), float(original), rel_tol=2e-15, abs_tol=1e-20)
            else:
                matches = actual == original
            if not matches:
                errors.append(f'{name}!{cell.coordinate}: source/readback mismatch')
    report_sheets.append({'sheet': name, 'data_rows': len(rows), 'columns': len(headers),
                          'source_csv': csv_data[name][0] if name in csv_data else None,
                          'source_csv_rows': len(csv_data[name][3]) if name in csv_data else None,
                          'row_count_matches': ws.max_row == len(rows) + 1, 'visible': ws.sheet_state == 'visible'})

archive = book['CSV_Literal_Archive']
for ridx, expected in enumerate(literal_rows, 2):
    recovered = json.loads(archive.cell(ridx, 3).value)
    source_path, source_line, _ = expected
    source_name = next(name for name, path, _ in CSV_INPUTS if path == source_path)
    _, _, headers, rows = csv_data[source_name]
    if recovered != dict(zip(headers, rows[source_line - 2])):
        errors.append(f'CSV literal archive mismatch at row {ridx}')
for path in JSON_INPUTS:
    recovered = ''.join(r[3] for r in json_literal_rows if r[0] == path)
    if recovered != (ROOT / path).read_text(encoding='utf-8'):
        errors.append('JSON literal archive mismatch: ' + path)
sha_after = {p: hashlib.sha256((ROOT / p).read_bytes()).hexdigest() for p in ALL_INPUTS}
if sha_after != SHA_BEFORE:
    errors.append('A frozen input changed while the workbook was being packaged')
if formula_cells:
    errors.append('Unexpected formula cells found')
book.close()
report = {'status': 'passed' if not errors else 'failed',
          'created_utc': dt.datetime.now(dt.timezone.utc).isoformat(),
          'workbook': str(DEST), 'workbook_sha256': hashlib.sha256(DEST.read_bytes()).hexdigest(),
          'workbook_bytes': DEST.stat().st_size, 'sheet_count': len(sheet_specs),
          'csv_sheet_count': len(CSV_INPUTS), 'csv_data_rows_total': sum(len(v[3]) for v in csv_data.values()),
          'all_frozen_input_sha256_unchanged': sha_after == SHA_BEFORE,
          'analysis_rerun': False, 'formula_cells': formula_cells,
          'literal_csv_archive_records': len(literal_rows), 'literal_json_chunks': len(json_literal_rows),
          'numeric_readback_relative_tolerance': 2e-15,
          'exact_original_csv_tokens_preserved': True,
          'numeric_storage_note': 'Excel numeric cells use double precision; all exact CSV tokens and original JSON text are also visibly archived as literal strings.',
          'source_input_sha256': SHA_BEFORE, 'sheets': report_sheets, 'errors': errors}
(OUT / 'excel_validation.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps({k: report[k] for k in ['status', 'workbook', 'workbook_bytes', 'sheet_count', 'csv_sheet_count', 'csv_data_rows_total', 'errors']}, ensure_ascii=False, indent=2))
if errors:
    raise SystemExit(1)
